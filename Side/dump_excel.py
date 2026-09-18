import openpyxl
import json
import os

excel_file = r"Side\RailBlock_Feature_Master_Plan_PS26027(1).xlsx"
wb = openpyxl.load_workbook(excel_file, data_only=True)

data = {}
for name in wb.sheetnames:
    ws = wb[name]
    rows = []
    headers = [str(ws.cell(row=1, column=c).value or '').strip() for c in range(1, ws.max_column+1)]
    for r in range(2, ws.max_row+1):
        row_dict = {}
        has_val = False
        for c in range(1, ws.max_column+1):
            val = ws.cell(row=r, column=c).value
            if val is not None and str(val).strip():
                has_val = True
            h_name = headers[c-1] if c-1 < len(headers) and headers[c-1] else f'col_{c}'
            row_dict[h_name] = val
        if has_val:
            rows.append(row_dict)
    data[name] = {'headers': headers, 'rows': rows}

out_dir = r"C:\Users\mrinm\.gemini\antigravity-ide\brain\b35975f5-cc5c-4d55-90cf-e0e124b2c600\scratch"
os.makedirs(out_dir, exist_ok=True)
out_path = os.path.join(out_dir, "excel_dump.json")

with open(out_path, "w", encoding="utf-8") as f:
    json.dump(data, f, ensure_ascii=False, indent=2)

print("SUCCESS: Dumped", len(data), "sheets to", out_path)
for k, v in data.items():
    print(f"Sheet '{k}': {len(v['rows'])} rows. Headers: {v['headers']}")
